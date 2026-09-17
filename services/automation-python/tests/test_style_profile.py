"""Tests for documents/style_profile.py.

The profile exists so a rebuilt resume can look like the one the user wrote
rather than like a generic template. Each test therefore builds a master with a
deliberately non-default look and asserts that look was read back, because a
profile that silently returns its own defaults would pass a weaker test while
producing a document the user does not recognise.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from applyocalypse_automation.documents.style_profile import (
    DEFAULT_STYLE_PROFILE,
    extract_docx_style_profile,
)


def _styled_run(paragraph, text: str, *, font: str, size: float, bold: bool = False):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    return run


def _write_master(path: Path) -> Path:
    """A resume that looks nothing like the builder's hardcoded defaults.

    Garamond 11.5 rather than Calibri 10, half-inch margins rather than
    three-quarter, a left-aligned name rather than centred, and the user's own
    section wording rather than "Skills"/"Experience".
    """
    document = Document()
    for section in document.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _styled_run(name, "Margaret Hamilton", font="Garamond", size=18, bold=True)

    contact = document.add_paragraph()
    _styled_run(contact, "margaret@mit.edu • 555-0100 • Cambridge, MA", font="Garamond", size=11.5)

    heading = document.add_paragraph()
    _styled_run(heading, "PROFESSIONAL EXPERIENCE", font="Garamond", size=12, bold=True)

    role = document.add_paragraph()
    _styled_run(role, "Lead Software Engineer, NASA", font="Garamond", size=11.5, bold=True)

    for text in (
        "Wrote the onboard flight software for the Apollo Guidance Computer.",
        "Designed priority scheduling that recovered the lander during overload.",
    ):
        bullet = document.add_paragraph(style="List Bullet")
        _styled_run(bullet, text, font="Garamond", size=11.5)

    education = document.add_paragraph()
    _styled_run(education, "EDUCATION", font="Garamond", size=12, bold=True)

    school = document.add_paragraph()
    _styled_run(school, "Earlham College, BA Mathematics", font="Garamond", size=11.5)

    document.save(str(path))
    return path


def test_reads_the_masters_margins(tmp_path: Path) -> None:
    profile = extract_docx_style_profile(_write_master(tmp_path / "master.docx"))
    assert profile.margins.left_in == pytest.approx(0.5)
    assert profile.margins.top_in == pytest.approx(0.5)


def test_reads_the_masters_body_font(tmp_path: Path) -> None:
    profile = extract_docx_style_profile(_write_master(tmp_path / "master.docx"))
    assert profile.body_font == "Garamond"
    # The dominant body size, not the larger name or heading runs.
    assert profile.body_size_pt == pytest.approx(11.5)


def test_keeps_the_users_own_section_wording_and_order(tmp_path: Path) -> None:
    # "PROFESSIONAL EXPERIENCE" must not come back as "Experience": renaming a
    # user's sections is the most visible way a rebuild stops being their resume.
    profile = extract_docx_style_profile(_write_master(tmp_path / "master.docx"))
    assert profile.section_labels == ("PROFESSIONAL EXPERIENCE", "EDUCATION")


def test_reads_the_name_line_treatment(tmp_path: Path) -> None:
    profile = extract_docx_style_profile(_write_master(tmp_path / "master.docx"))
    assert profile.name_size_pt == pytest.approx(18)
    assert profile.name_bold is True
    assert profile.name_centered is False


def test_reads_the_contact_separator(tmp_path: Path) -> None:
    profile = extract_docx_style_profile(_write_master(tmp_path / "master.docx"))
    assert profile.contact_separator.strip() == "•"


def test_reads_the_bullet_style(tmp_path: Path) -> None:
    profile = extract_docx_style_profile(_write_master(tmp_path / "master.docx"))
    assert profile.bullet_style == "List Bullet"


def test_notices_all_caps_headings(tmp_path: Path) -> None:
    profile = extract_docx_style_profile(_write_master(tmp_path / "master.docx"))
    assert profile.heading_all_caps is True
    assert profile.heading_bold is True


def test_centred_name_is_reported_as_centred(tmp_path: Path) -> None:
    path = tmp_path / "centred.docx"
    document = Document()
    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _styled_run(name, "Grace Hopper", font="Calibri", size=16, bold=True)
    body = document.add_paragraph()
    _styled_run(body, "grace@navy.mil | 555-0111", font="Calibri", size=10)
    document.save(str(path))

    assert extract_docx_style_profile(path).name_centered is True


def test_unreadable_master_falls_back_and_says_so(tmp_path: Path) -> None:
    # A profile is a nicety, so a broken master must not stop a run. It must
    # also not claim it read a style it never saw.
    broken = tmp_path / "not-really.docx"
    broken.write_text("this is not a docx", encoding="utf-8")

    profile = extract_docx_style_profile(broken)

    assert profile.detected is False
    assert profile.body_font == DEFAULT_STYLE_PROFILE.body_font


def test_a_read_profile_reports_itself_as_detected(tmp_path: Path) -> None:
    assert extract_docx_style_profile(_write_master(tmp_path / "master.docx")).detected is True
