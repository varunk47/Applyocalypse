from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
import re
import shutil

from .docx_mutation import iter_document_paragraphs, replace_paragraph_text_preserving_runs


REQUIRED_PLACEHOLDERS = (
    "{{APPLYO_FULL_NAME}}",
    "{{APPLYO_RESUME_SUMMARY}}",
    "{{APPLYO_SKILLS}}",
    "{{APPLYO_EXP_0_BULLETS}}",
    "{{APPLYO_EXP_1_BULLETS}}",
    "{{APPLYO_EXP_2_BULLETS}}",
    "{{APPLYO_EXP_3_BULLETS}}",
    "{{APPLYO_PROJECTS_BULLETS}}",
)

SUMMARY_LABELS = {"summary", "profile", "professional summary", "about me", "objective"}
SKILLS_LABELS = {"skills", "technical skills", "technologies", "tools", "core skills", "competencies"}
EXPERIENCE_LABELS = {"experience", "work experience", "employment", "professional experience", "career history"}
PROJECTS_LABELS = {"projects", "personal projects", "key projects", "selected projects", "academic projects"}


@dataclass(frozen=True, slots=True)
class AnchorRepairResult:
    source_path: str
    output_path: str
    source_format: str
    added_placeholders: list[str]
    already_present_placeholders: list[str]
    warnings: list[str]

    def to_payload(self) -> dict[str, object]:
        return asdict(self)


def normalize_label(value: str) -> str:
    cleaned = re.sub(r"[^a-z0-9 &/+-]", "", value.strip().strip(":").lower())
    return re.sub(r"\s+", " ", cleaned)


def detect_placeholders(text: str) -> list[str]:
    return [placeholder for placeholder in REQUIRED_PLACEHOLDERS if placeholder in text]


def is_likely_contact_line(value: str) -> bool:
    lowered = value.lower()
    return "@" in value or "http" in lowered or "linkedin" in lowered or bool(re.search(r"\+?\d[\d\s().-]{7,}\d", value))


def first_candidate_name_paragraph(paragraphs: list[object]) -> object | None:
    for paragraph in paragraphs[:8]:
        text = str(getattr(paragraph, "text", "")).strip()
        if not text or is_likely_contact_line(text):
            continue
        if 2 <= len(text.split()) <= 6 and len(text) <= 90:
            return paragraph
    return None


def next_non_empty_paragraph(paragraphs: list[object], start_index: int) -> object | None:
    for paragraph in paragraphs[start_index + 1 :]:
        text = str(getattr(paragraph, "text", "")).strip()
        if text:
            return paragraph
    return None


def find_body_after_section(paragraphs: list[object], labels: set[str]) -> object | None:
    for index, paragraph in enumerate(paragraphs):
        if normalize_label(str(getattr(paragraph, "text", ""))) in labels:
            return next_non_empty_paragraph(paragraphs, index)
    return None


def find_first_bullet_in_each_experience_block(paragraphs: list[object]) -> list[object]:
    """Return the first bullet paragraph of each work experience entry.

    Heuristic: find the experience section heading, then detect sub-headings
    (company/role lines) followed by bullet paragraphs. Returns one bullet per entry.
    """
    exp_section_idx = None
    for index, paragraph in enumerate(paragraphs):
        if normalize_label(str(getattr(paragraph, "text", ""))) in EXPERIENCE_LABELS:
            exp_section_idx = index
            break
    if exp_section_idx is None:
        return []

    first_bullets: list[object] = []
    in_exp_section = False
    last_sub_heading_idx = None

    for index, paragraph in enumerate(paragraphs):
        if index <= exp_section_idx:
            if index == exp_section_idx:
                in_exp_section = True
            continue
        text = str(getattr(paragraph, "text", "")).strip()
        if not text:
            continue
        normalized = normalize_label(text)
        # Stop when we hit another section heading of similar weight
        if normalized in SKILLS_LABELS or normalized in SUMMARY_LABELS or normalized in PROJECTS_LABELS:
            break
        # Detect sub-headings (company name / role lines): short, not bullet style
        is_bullet = (
            getattr(paragraph, "_p", None) is not None
            and getattr(getattr(paragraph, "_p", None), "pPr", None) is not None
            and getattr(getattr(paragraph, "_p", None).pPr, "numPr", None) is not None
        )
        is_sub_heading = not is_bullet and 1 <= len(text.split()) <= 8
        if is_sub_heading and in_exp_section:
            last_sub_heading_idx = index
            continue
        if is_bullet and last_sub_heading_idx is not None:
            first_bullets.append(paragraph)
            last_sub_heading_idx = None  # consume — only take the first bullet per entry

    return first_bullets


def find_first_bullet_in_projects(paragraphs: list[object]) -> object | None:
    return find_body_after_section(paragraphs, PROJECTS_LABELS)


def repair_docx_anchors(source: Path, output: Path) -> AnchorRepairResult:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX anchor repair") from exc

    document = Document(str(source))
    paragraphs = iter_document_paragraphs(document)
    full_text = "\n".join(paragraph.text for paragraph in paragraphs)
    already_present = detect_placeholders(full_text)
    added: list[str] = []
    warnings: list[str] = []

    if "{{APPLYO_FULL_NAME}}" not in already_present:
        paragraph = first_candidate_name_paragraph(paragraphs)
        if paragraph is None:
            warnings.append("Could not identify a high-confidence legal-name paragraph.")
        else:
            replace_paragraph_text_preserving_runs(paragraph, "{{APPLYO_FULL_NAME}}")
            added.append("{{APPLYO_FULL_NAME}}")

    if "{{APPLYO_RESUME_SUMMARY}}" not in already_present:
        paragraph = find_body_after_section(paragraphs, SUMMARY_LABELS)
        if paragraph is None:
            warnings.append("No summary/profile section body found for {{APPLYO_RESUME_SUMMARY}}.")
        else:
            replace_paragraph_text_preserving_runs(paragraph, "{{APPLYO_RESUME_SUMMARY}}")
            added.append("{{APPLYO_RESUME_SUMMARY}}")

    if "{{APPLYO_SKILLS}}" not in already_present:
        paragraph = find_body_after_section(paragraphs, SKILLS_LABELS)
        if paragraph is None:
            warnings.append("No skills section body found for {{APPLYO_SKILLS}}.")
        else:
            replace_paragraph_text_preserving_runs(paragraph, "{{APPLYO_SKILLS}}")
            added.append("{{APPLYO_SKILLS}}")

    exp_anchors = [f"{{{{APPLYO_EXP_{i}_BULLETS}}}}" for i in range(4)]
    missing_exp = [a for a in exp_anchors if a not in already_present]
    if missing_exp:
        exp_bullets = find_first_bullet_in_each_experience_block(paragraphs)
        for anchor, paragraph in zip(missing_exp, exp_bullets):
            replace_paragraph_text_preserving_runs(paragraph, anchor)
            added.append(anchor)
        if not exp_bullets:
            warnings.append("No work experience bullet blocks found for {{APPLYO_EXP_N_BULLETS}} anchors.")

    if "{{APPLYO_PROJECTS_BULLETS}}" not in already_present:
        paragraph = find_first_bullet_in_projects(paragraphs)
        if paragraph is None:
            warnings.append("No projects section found for {{APPLYO_PROJECTS_BULLETS}}.")
        else:
            replace_paragraph_text_preserving_runs(paragraph, "{{APPLYO_PROJECTS_BULLETS}}")
            added.append("{{APPLYO_PROJECTS_BULLETS}}")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return AnchorRepairResult(
        source_path=str(source),
        output_path=str(output),
        source_format="DOCX",
        added_placeholders=added,
        already_present_placeholders=already_present,
        warnings=warnings,
    )


def replace_tex_section_body(source: str, labels: set[str], placeholder: str) -> tuple[str, bool]:
    section_pattern = re.compile(r"(\\section\*?\{([^}]*)\})(.*?)(?=\\section\*?\{|\\end\{document\}|\Z)", re.DOTALL)
    for match in section_pattern.finditer(source):
        label = normalize_label(match.group(2))
        if label not in labels:
            continue
        replacement = f"{match.group(1)}\n{placeholder}\n"
        return f"{source[:match.start()]}{replacement}{source[match.end():]}", True
    return source, False


def repair_tex_anchors(source: Path, output: Path) -> AnchorRepairResult:
    text = source.read_text(encoding="utf-8")
    already_present = detect_placeholders(text)
    added: list[str] = []
    warnings: list[str] = []
    repaired = text

    if "{{APPLYO_FULL_NAME}}" not in already_present:
        warnings.append("TEX name/title macros are not auto-repaired because bare placeholders inside command braces can corrupt LaTeX syntax.")

    if "{{APPLYO_RESUME_SUMMARY}}" not in already_present:
        repaired, replaced = replace_tex_section_body(repaired, SUMMARY_LABELS, "{{APPLYO_RESUME_SUMMARY}}")
        if replaced:
            added.append("{{APPLYO_RESUME_SUMMARY}}")
        else:
            warnings.append("No summary/profile section body found for {{APPLYO_RESUME_SUMMARY}}.")

    if "{{APPLYO_SKILLS}}" not in already_present:
        repaired, replaced = replace_tex_section_body(repaired, SKILLS_LABELS, "{{APPLYO_SKILLS}}")
        if replaced:
            added.append("{{APPLYO_SKILLS}}")
        else:
            warnings.append("No skills section body found for {{APPLYO_SKILLS}}.")

    if not added and not already_present:
        insertion = "\n\\section*{Applyocalypse Review Anchors}\n{{APPLYO_RESUME_SUMMARY}}\n{{APPLYO_SKILLS}}\n"
        if "\\end{document}" in repaired:
            repaired = repaired.replace("\\end{document}", f"{insertion}\\end{{document}}", 1)
            added.extend(["{{APPLYO_RESUME_SUMMARY}}", "{{APPLYO_SKILLS}}"])
            warnings.append("Inserted a review-only anchor section. Move these placeholders into the intended layout before confirmation.")

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(repaired, encoding="utf-8")
    return AnchorRepairResult(
        source_path=str(source),
        output_path=str(output),
        source_format="TEX",
        added_placeholders=added,
        already_present_placeholders=already_present,
        warnings=warnings,
    )


def repair_editable_master_anchors(source: Path, output: Path) -> AnchorRepairResult:
    suffix = source.suffix.lower()
    if suffix == ".docx":
        return repair_docx_anchors(source, output)
    if suffix == ".tex":
        return repair_tex_anchors(source, output)
    if suffix in {".txt", ".md"}:
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, output)
        return AnchorRepairResult(
            source_path=str(source),
            output_path=str(output),
            source_format=suffix[1:].upper(),
            added_placeholders=[],
            already_present_placeholders=[],
            warnings=["Anchor repair is only available for DOCX and TEX editable masters."],
        )
    raise ValueError("Anchor repair is only available for DOCX and TEX editable masters.")
